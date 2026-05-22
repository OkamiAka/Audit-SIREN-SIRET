import os
import re
import time
import uuid
import tempfile
import logging
import unicodedata
from dataclasses import dataclass
from difflib import SequenceMatcher

import pandas as pd
import requests
from flask import Flask, render_template, request, redirect, url_for, flash, send_file

from openpyxl import load_workbook
from openpyxl.styles import PatternFill

# =========================
# Config
# =========================
BASE_URL = "https://api.insee.fr/api-sirene/3.11"  # 【4-1ba1b7】
INSEE_API_KEY = os.getenv("INSEE_API_KEY", "")
HEADERS = {"X-INSEE-Api-Key-Integration": INSEE_API_KEY} if INSEE_API_KEY else {}  # 【4-2d78a6】

MAX_REQ_PER_MIN = 28  # < 30/min (safe) 【1-7defd0】【2-87bf3d】
MIN_INTERVAL = 60.0 / MAX_REQ_PER_MIN

REQUIRED = [
    "Client", "NomClient", "NafClient", "SIRET",
    "CPostalClient", "VilleClient", "Adresse1Client", "Adresse2Client"
]  # 【3-f02937】

LOG_DIR = os.path.join(os.getcwd(), "logs")
os.makedirs(LOG_DIR, exist_ok=True)
LOG_PATH = os.path.join(LOG_DIR, "app.log")

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s | %(levelname)s | %(message)s",
    handlers=[logging.FileHandler(LOG_PATH, encoding="utf-8"), logging.StreamHandler()],
)

app = Flask(__name__)
app.secret_key = os.getenv("FLASK_SECRET", "change-me")

# =========================
# Rate-limited HTTP session
# =========================
session = requests.Session()
_last_call_ts = 0.0

def _rate_limit_wait():
    global _last_call_ts
    now = time.time()
    wait = MIN_INTERVAL - (now - _last_call_ts)
    if wait > 0:
        time.sleep(wait)
    _last_call_ts = time.time()

def http_get(url, params=None):
    """
    Strict <30/min + retry on 429.
    """
    for attempt in range(6):
        _rate_limit_wait()
        r = session.get(url, headers=HEADERS, params=params, timeout=25)
        if r.status_code == 429:
            sleep_s = 2.5 + attempt * 2.5
            logging.warning("429 Too Many Requests -> backoff %.1fs", sleep_s)
            time.sleep(sleep_s)
            continue
        return r
    return r

# =========================
# Small caches (avoid extra calls)
# =========================
CACHE_MAX = 5000
cache_siren = {}   # siren -> uniteLegale
cache_siret = {}   # siret -> etablissement (or None)
cache_q = {}       # (q,limit) -> etablissements list

def _cache_put(d, key, value):
    if key in d:
        d[key] = value
        return
    if len(d) >= CACHE_MAX:
        # simple eviction: pop one item
        d.pop(next(iter(d)))
    d[key] = value

# =========================
# Helpers
# =========================
def strip_accents(s: str) -> str:
    s = str(s or "")
    return "".join(
        c for c in unicodedata.normalize("NFD", s)
        if unicodedata.category(c) != "Mn"
    )

def digits_only(x) -> str:
    return re.sub(r"\D", "", str(x or ""))

def normalize_id(x: str) -> str:
    s = digits_only(x)
    if len(s) == 13:  # Excel a mangé un 0
        s = s.zfill(14)
    return s

def clean_for_match(s: str) -> str:
    s = strip_accents(str(s or "")).upper()
    s = re.sub(r"[\"'()]", " ", s)
    s = re.sub(r"\s+", " ", s).strip()
    return s

def sim(a: str, b: str) -> float:
    return SequenceMatcher(None, clean_for_match(a), clean_for_match(b)).ratio()

def norm_naf(x: str) -> str:
    s = strip_accents(str(x or "")).upper()
    return s.replace(".", "").replace(" ", "")

def build_nom_officiel(ul: dict) -> str:
    if not isinstance(ul, dict):
        return ""
    for k in [
        "denominationUniteLegale",
        "denominationUsuelle1UniteLegale",
        "denominationUsuelle2UniteLegale",
        "denominationUsuelle3UniteLegale",
        "sigleUniteLegale",
    ]:
        v = (ul.get(k) or "").strip()
        if v:
            return v
    nom = (ul.get("nomUniteLegale") or "").strip()
    prenom1 = (ul.get("prenom1UniteLegale") or "").strip()
    prenom_usuel = (ul.get("prenomUsuelUniteLegale") or "").strip()
    nom_usage = (ul.get("nomUsageUniteLegale") or "").strip()
    if nom and (prenom_usuel or prenom1):
        return f"{nom} {prenom_usuel or prenom1}".strip()
    if nom_usage and (prenom_usuel or prenom1):
        return f"{nom_usage} {prenom_usuel or prenom1}".strip()
    return nom or ""

def build_adresse_api(adresse: dict) -> str:
    if not isinstance(adresse, dict):
        return ""
    parts = []
    for k in [
        "complementAdresseEtablissement",
        "numeroVoieEtablissement",
        "indiceRepetitionEtablissement",
        "typeVoieEtablissement",
        "libelleVoieEtablissement",
        "distributionSpecialeEtablissement",
    ]:
        v = (adresse.get(k) or "").strip()
        if v:
            parts.append(v)
    return " ".join(parts).strip()

def etat_admin_etab(e: dict) -> str:
    if not isinstance(e, dict):
        return ""
    v = (e.get("etatAdministratifEtablissement") or "").strip()
    if v:
        return v
    periodes = e.get("periodesEtablissement") or []
    if isinstance(periodes, list) and periodes:
        current = None
        for p in periodes:
            if isinstance(p, dict) and not p.get("dateFin"):
                current = p
                break
        p = current or periodes[0]
        return (p.get("etatAdministratifEtablissement") or "").strip()
    return ""

def is_actif(e: dict) -> bool:
    return etat_admin_etab(e) == "A"

# =========================
# Fix colonnes décalées (cas réel) 【3-f02937】
# NafClient=SIREN, SIRET=CP, CPostalClient=Ville
# =========================
def fix_row(row: dict) -> dict:
    row = dict(row)

    naf = str(row.get("NafClient", "") or "").strip()
    siret = str(row.get("SIRET", "") or "").strip()
    cp = str(row.get("CPostalClient", "") or "").strip()
    ville = str(row.get("VilleClient", "") or "").strip()

    naf_d = digits_only(naf)
    siret_d = digits_only(siret)
    looks_like_city_in_cp = any(c.isalpha() for c in cp)

    if len(naf_d) == 9 and len(siret_d) == 5 and (looks_like_city_in_cp or not ville):
        row["SIREN_FIX"] = naf_d
        row["CPostalClient"] = siret_d
        row["VilleClient"] = cp
        row["SIRET"] = ""
        row["NafClient"] = ""
        return row

    if len(naf_d) == 9:
        row["SIREN_FIX"] = naf_d
        row["NafClient"] = ""
        return row

    row["SIREN_FIX"] = ""
    return row

# =========================
# API wrappers
# =========================
def api_get_ul_by_siren(siren9: str):
    if siren9 in cache_siren:
        return cache_siren[siren9], 200

    url = f"{BASE_URL}/siren/{siren9}"  # 【4-4dcb19】
    r = http_get(url)
    if r.status_code != 200:
        return None, r.status_code
    js = r.json()
    ul = js.get("uniteLegale", None)
    _cache_put(cache_siren, siren9, ul)
    return ul, 200

def api_get_etab_by_siret(siret14: str):
    siret14 = normalize_id(siret14)
    if len(siret14) != 14:
        return None, 400
    if siret14 in cache_siret:
        return cache_siret[siret14], 200

    url = f"{BASE_URL}/siret/{siret14}"  # 【4-0e67f0】
    r = http_get(url)
    if r.status_code != 200:
        _cache_put(cache_siret, siret14, None)
        return None, r.status_code
    js = r.json()
    etab = js.get("etablissement", None)
    _cache_put(cache_siret, siret14, etab)
    return etab, 200

def api_search_etabs(q: str, limit=200):
    key = (q, limit)
    if key in cache_q:
        return cache_q[key], 200

    url = f"{BASE_URL}/siret"
    params = {"q": q, "nombre": limit, "curseur": "*"}  # 【4-b804f0】
    r = http_get(url, params=params)
    if r.status_code == 200:
        etabs = r.json().get("etablissements", [])
        _cache_put(cache_q, key, etabs)
        return etabs, 200
    if r.status_code == 404:
        _cache_put(cache_q, key, [])
        return [], 200
    return [], r.status_code

# =========================
# SIREN -> NIC siège
# =========================
def extract_nic_siege(ul: dict) -> str:
    if not isinstance(ul, dict):
        return ""
    periodes = ul.get("periodesUniteLegale") or []
    if isinstance(periodes, list) and periodes:
        current = None
        for p in periodes:
            if isinstance(p, dict) and p.get("dateFin") is None:
                current = p
                break
        p = current or periodes[0]
        nic = (p.get("nicSiegeUniteLegale") or "").strip()
        return nic
    return ""

# =========================
# Choose best establishment
# =========================
def choose_best_etab(etabs: list, cp: str, ville: str, nom: str):
    if not etabs:
        return None, 0.0, False

    cp_d = digits_only(cp)
    has_active = any(is_actif(e) for e in etabs)
    best, best_score = None, -1e9

    for e in etabs:
        ul = e.get("uniteLegale", {}) or {}
        nom_off = build_nom_officiel(ul)

        adr = e.get("adresseEtablissement", {}) or {}
        cp_api = (adr.get("codePostalEtablissement") or "").strip()
        ville_api = (adr.get("libelleCommuneEtablissement") or "").strip()

        s_nom = sim(nom, nom_off) if nom_off else 0.0
        s_ville = sim(ville, ville_api) if ville and ville_api else 0.0
        s_cp = 1.0 if (cp_d and cp_d == digits_only(cp_api)) else 0.0

        score = 0.70 * s_nom + 0.20 * s_ville + 0.10 * s_cp
        if e.get("etablissementSiege") is True:
            score += 0.05

        if has_active and not is_actif(e):
            score -= 0.50
        elif (not has_active) and not is_actif(e):
            score -= 0.10

        if score > best_score:
            best_score, best = score, e

    return best, round(best_score, 2), has_active

# =========================
# Fallback name -> siret
# =========================
def find_siret_by_name(nom: str, cp: str, ville: str):
    nom_clean = clean_for_match(nom)
    cp_d = digits_only(cp)

    queries = [f'denominationUniteLegale:"{nom_clean}"']
    if cp_d:
        queries.insert(0, f'periode(etatAdministratifEtablissement:A) AND denominationUniteLegale:"{nom_clean}" AND codePostalEtablissement:{cp_d}')
    else:
        queries.insert(0, f'periode(etatAdministratifEtablissement:A) AND denominationUniteLegale:"{nom_clean}"')

    for q in queries:
        etabs, st = api_search_etabs(q, limit=200)
        if st != 200:
            continue
        if etabs:
            best, _, _ = choose_best_etab(etabs, cp, ville, nom)
            if best:
                siret = normalize_id(best.get("siret", "") or "")
                if len(siret) == 14:
                    return siret
    return ""

# =========================
# Enrich output via SIRET
# =========================
def enrich_from_siret(out: dict, siret14: str):
    siret14 = normalize_id(siret14)
    etab, st = api_get_etab_by_siret(siret14)
    if st != 200 or not etab:
        return out

    ul = etab.get("uniteLegale", {}) or {}
    nom_off = build_nom_officiel(ul)

    adr = etab.get("adresseEtablissement", {}) or {}
    cp_api = (adr.get("codePostalEtablissement") or "").strip()
    ville_api = (adr.get("libelleCommuneEtablissement") or "").strip()
    adresse_api = build_adresse_api(adr)

    naf_api = etab.get("activitePrincipaleEtablissement", "") or ul.get("activitePrincipaleUniteLegale", "") or ""
    naf_api = norm_naf(naf_api)

    out.update({
        "EtatEtab_API": etat_admin_etab(etab),
        "SIREN_API": etab.get("siren", siret14[:9]) or siret14[:9],
        "SIRET_API": siret14,
        "Nom_API": nom_off,
        "NAF_API": naf_api,
        "CP_API": cp_api,
        "Ville_API": ville_api,
        "Adresse_API": adresse_api,
        "Score": round(sim(out.get("NomClient",""), nom_off), 2) if nom_off else 0.0
    })
    return out

# =========================
# SIREN -> SIRET siège actif
# =========================
def resolve_from_siren(siren9: str, nom: str, cp: str, ville: str):
    ul, st = api_get_ul_by_siren(siren9)
    if st != 200 or not ul:
        return "", "SIREN introuvable"

    nic = extract_nic_siege(ul)
    if not nic or len(nic) != 5:
        return "", "NIC siège introuvable"

    siret_siege = siren9 + nic
    etab_siege, st2 = api_get_etab_by_siret(siret_siege)

    if st2 == 200 and etab_siege:
        if is_actif(etab_siege):
            return siret_siege, "OK (siège actif)"
        # siège fermé -> chercher un actif
        qA = f"periode(etatAdministratifEtablissement:A) AND periode(siren:{siren9})"
        etabs, st3 = api_search_etabs(qA, limit=200)
        if st3 == 200 and etabs:
            best, _, _ = choose_best_etab(etabs, cp, ville, nom)
            if best:
                siret = normalize_id(best.get("siret", "") or "")
                if len(siret) == 14:
                    return siret, f"⚠️ Siège fermé → établissement actif choisi ({siret})"
        return siret_siege, "⚠️ Siège fermé et aucun établissement actif trouvé"

    # fallback : liste active
    qA = f"periode(etatAdministratifEtablissement:A) AND periode(siren:{siren9})"
    etabs, st3 = api_search_etabs(qA, limit=200)
    if st3 == 200 and etabs:
        best, _, _ = choose_best_etab(etabs, cp, ville, nom)
        if best:
            siret = normalize_id(best.get("siret", "") or "")
            if len(siret) == 14:
                return siret, "OK (actif trouvé)"
    return "", "Impossible de récupérer un SIRET via SIREN"

# =========================
# Business status
# =========================
def compute_business_status(out: dict) -> str:
    problems = []
    nom_client = out.get("NomClient","")
    nom_api = out.get("Nom_API","")
    cp_client = digits_only(out.get("CPostalClient",""))
    cp_api = digits_only(out.get("CP_API",""))
    ville_client = out.get("VilleClient","")
    ville_api = out.get("Ville_API","")
    naf_client = norm_naf(out.get("NafClient",""))
    naf_api = norm_naf(out.get("NAF_API",""))

    if nom_api and sim(nom_client, nom_api) < 0.6:
        problems.append(f"NOM différent → {nom_api}")
    if cp_client and cp_api and cp_client != cp_api:
        problems.append(f"CP différent → {cp_api}")
    if ville_client and ville_api and sim(ville_client, ville_api) < 0.6:
        problems.append(f"VILLE différente → {ville_api}")
    if naf_client and naf_api and naf_client != naf_api:
        problems.append(f"NAF différent → {naf_api}")

    return "OK ✅" if not problems else "⚠️ " + " | ".join(problems)

# =========================
# Core verify
# =========================
def verify_line(row: dict) -> dict:
    row = fix_row(row)

    client = row.get("Client","")
    nom = row.get("NomClient","")
    naf = row.get("NafClient","")
    siret_excel = row.get("SIRET","")
    cp = row.get("CPostalClient","")
    ville = row.get("VilleClient","")

    id_norm = normalize_id(siret_excel)
    siren_fix = str(row.get("SIREN_FIX","") or "").strip()
    if (not id_norm) and siren_fix and siren_fix.isdigit() and len(siren_fix) == 9:
        id_norm = siren_fix

    out = {
        "Client": client,
        "NomClient": nom,
        "NafClient": naf,
        "SIRET_Excel": siret_excel,
        "CPostalClient": cp,
        "VilleClient": ville,
        "Type_ID": "",
        "EtatEtab_API": "",
        "SIREN_API": "",
        "SIRET_API": "",
        "Nom_API": "",
        "NAF_API": "",
        "CP_API": "",
        "Ville_API": "",
        "Adresse_API": "",
        "Score": 0.0,
        "Statut": "",
    }

    # SIRET direct
    if id_norm.isdigit() and len(id_norm) == 14:
        out["Type_ID"] = "SIRET"
        out["Statut"] = "OK (SIRET fourni) | "
        out = enrich_from_siret(out, id_norm)
        out["Statut"] += compute_business_status(out)
        return out

    # SIREN -> SIRET siège actif
    if id_norm.isdigit() and len(id_norm) == 9:
        out["Type_ID"] = "SIREN"
        siret_found, msg = resolve_from_siren(id_norm, nom, cp, ville)
        if siret_found:
            out["Statut"] = msg + " | "
            out = enrich_from_siret(out, siret_found)
            out["Statut"] += compute_business_status(out)
            return out

        # fallback NOM
        siret_by_name = find_siret_by_name(nom, cp, ville)
        if siret_by_name:
            out["Type_ID"] = "SIREN (fallback NOM)"
            out["Statut"] = msg + " → trouvé via NOM | "
            out = enrich_from_siret(out, siret_by_name)
            out["Statut"] += compute_business_status(out)
            return out

        out["Statut"] = msg + " + NOM introuvable"
        return out

    # fallback NOM
    out["Type_ID"] = "NOM"
    siret_by_name = find_siret_by_name(nom, cp, ville)
    if siret_by_name:
        out["Statut"] = "OK (trouvé via NOM) | "
        out = enrich_from_siret(out, siret_by_name)
        out["Statut"] += compute_business_status(out)
        return out

    out["Statut"] = "Non trouvé"
    return out

# =========================
# Excel color export (server-side)
# =========================
FILL_OK = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")     # green
FILL_WARN = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")  # yellow
FILL_ERR = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")   # red

def colorize_excel(path_xlsx: str):
    wb = load_workbook(path_xlsx)
    ws = wb.active

    # find "Statut" col index
    header = [cell.value for cell in ws[1]]
    try:
        statut_idx = header.index("Statut") + 1
    except ValueError:
        wb.save(path_xlsx)
        return

    for r in range(2, ws.max_row + 1):
        statut = ws.cell(row=r, column=statut_idx).value or ""
        if "OK ✅" in str(statut) or str(statut).startswith("OK"):
            fill = FILL_OK
        elif "⚠️" in str(statut):
            fill = FILL_WARN
        else:
            fill = FILL_ERR
        for c in range(1, ws.max_column + 1):
            ws.cell(row=r, column=c).fill = fill

    wb.save(path_xlsx)

# =========================
# Flask routes
# =========================
@app.route("/", methods=["GET"])
def index():
    return render_template("index.html", rows=None, download_id=None)

@app.route("/health", methods=["GET"])
def health():
    return {"status": "ok", "rate_limit_per_min": MAX_REQ_PER_MIN}, 200

@app.route("/process", methods=["POST"])
def process():
    f = request.files.get("file")
    if not f or not f.filename:
        flash("Importe ton Excel BASE CLT INFO.")
        return redirect(url_for("index"))

    tmp_in = os.path.join(tempfile.gettempdir(), f"upload_{uuid.uuid4().hex}.xlsx")
    f.save(tmp_in)

    df = pd.read_excel(tmp_in, engine="openpyxl", dtype=str)

    missing = [c for c in REQUIRED if c not in df.columns]
    if missing:
        flash(f"Colonnes manquantes: {', '.join(missing)}")
        return redirect(url_for("index"))

    results = []
    for _, r in df.iterrows():
        row = {k: (r.get(k) if pd.notna(r.get(k)) else "") for k in REQUIRED}
        results.append(verify_line(row))

    out_df = pd.DataFrame(results)
    out_id = uuid.uuid4().hex[:10]
    tmp_out = os.path.join(tempfile.gettempdir(), f"result_{out_id}.xlsx")
    out_df.to_excel(tmp_out, index=False, engine="openpyxl")

    # colorize server-side
    colorize_excel(tmp_out)

    logging.info("Export généré: %s (lignes=%d)", tmp_out, len(out_df))

    return render_template("index.html", rows=results, download_id=out_id)

@app.route("/download/<download_id>", methods=["GET"])
def download(download_id):
    path = os.path.join(tempfile.gettempdir(), f"result_{download_id}.xlsx")
    if not os.path.exists(path):
        flash("Résultat introuvable.")
        return redirect(url_for("index"))
    return send_file(path, as_attachment=True, download_name="result_verif_base_clt.xlsx")

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000, debug=True)
